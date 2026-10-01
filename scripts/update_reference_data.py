from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import date
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DOWNLOAD_DIR = ROOT / ".reference-downloads"
OUTPUT = ROOT / "services" / "api" / "app" / "reference_data" / "catalog.json"

DEFAULT_VALUES_FILE = DOWNLOAD_DIR / "cbam-default-values-2026-08-10.xlsx"
BENCHMARKS_FILE = DOWNLOAD_DIR / "cbam-benchmarks-2026-02-13.xlsx"

DEFAULT_VALUES_URL = (
    "https://taxation-customs.ec.europa.eu/document/download/"
    "1c05d211-80cb-4aaa-8ef0-e08005a95d7e_en"
    "?filename=DV+correcting+act_final+update_06.08.xlsx"
)
BENCHMARKS_URL = (
    "https://taxation-customs.ec.europa.eu/document/download/"
    "9877523c-2a02-4926-a211-aefae7cf6d0d_en"
    "?filename=CBAM+Benchmarks_20260206.xlsx"
)

ALLOWED_SOURCE_HOSTS = {"taxation-customs.ec.europa.eu"}
MAX_DOWNLOAD_BYTES = 50 * 1024 * 1024

SECTOR_NAMES = {
    "Cement": "Cement",
    "Fertilisers": "Fertilisers",
    "Iron and steel": "Iron and steel",
    "Iron & Steel": "Iron and steel",
    "Aluminium": "Aluminium",
    "Hydrogen": "Hydrogen",
    "Electricity": "Electricity",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_source_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in ALLOWED_SOURCE_HOSTS:
        raise ValueError(f"Unapproved reference data URL: {url}")


def download_file(url: str, destination: Path) -> None:
    validate_source_url(url)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(f"{destination.suffix}.part")
    request = Request(
        url,
        headers={"User-Agent": "Marj reference data updater/0.2"},
    )

    try:
        with urlopen(request, timeout=60) as response, temporary.open("wb") as target:
            validate_source_url(response.geturl())
            total = 0
            while chunk := response.read(1024 * 1024):
                total += len(chunk)
                if total > MAX_DOWNLOAD_BYTES:
                    raise ValueError(
                        f"Reference file exceeds {MAX_DOWNLOAD_BYTES} bytes: {url}"
                    )
                target.write(chunk)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)


def clean_cn_code(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    digits = re.sub(r"\D", "", str(value))
    return digits or None


def parse_decimal(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    normalized = str(value).strip().replace(" ", "").replace(",", ".")
    if normalized in {"", "-", "N/A", "n/a", "seebelow"}:
        return None
    try:
        return float(normalized)
    except ValueError:
        return None


def normalize_route(value: Any) -> str | None:
    if value is None:
        return None
    route = str(value).strip()
    return route if route else None


def find_turkiye_sheet(workbook: Any) -> str:
    for sheet_name in workbook.sheetnames:
        if sheet_name.casefold() in {"türkiye", "turkiye"}:
            return sheet_name
    raise RuntimeError("Türkiye sheet was not found in the default values workbook")


def parse_default_values(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[find_turkiye_sheet(workbook)]
    sector: str | None = None
    rows: list[dict[str, Any]] = []

    for row in sheet.iter_rows(min_row=3, values_only=True):
        label = str(row[0]).strip() if row[0] is not None else ""
        if label in SECTOR_NAMES:
            sector = SECTOR_NAMES[label]
            continue

        cn_code = clean_cn_code(row[0])
        total = parse_decimal(row[4])
        direct = parse_decimal(row[2])
        indirect = parse_decimal(row[3])
        if not sector or not cn_code or (total is None and direct is None):
            continue

        rows.append(
            {
                "origin_country": "Türkiye",
                "sector": sector,
                "cn_code": cn_code,
                "description": str(row[1]).strip(),
                "direct_tco2e_per_tonne": direct,
                "indirect_tco2e_per_tonne": indirect,
                "total_tco2e_per_tonne": total if total is not None else direct,
                "production_route": normalize_route(row[5]),
            }
        )

    return rows


def parse_benchmarks(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Benchmarks"]
    sector: str | None = None
    current_cn: str | None = None
    current_description: str | None = None
    rows: list[dict[str, Any]] = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        label = str(row[0]).strip() if row[0] is not None else ""
        if label in SECTOR_NAMES:
            sector = SECTOR_NAMES[label]
            current_cn = None
            current_description = None
            continue

        candidate_cn = clean_cn_code(row[0])
        if candidate_cn:
            current_cn = candidate_cn
            current_description = str(row[1]).strip()

        column_a = parse_decimal(row[2])
        column_b = parse_decimal(row[4])
        if not sector or not current_cn or (column_a is None and column_b is None):
            continue

        rows.append(
            {
                "sector": sector,
                "cn_code": current_cn,
                "description": current_description,
                "column_a_tco2e_per_tonne": column_a,
                "column_a_route": normalize_route(row[3]),
                "column_b_tco2e_per_tonne": column_b,
                "column_b_route": normalize_route(row[5]),
            }
        )

    return rows


def build_catalog() -> dict[str, Any]:
    for path in (DEFAULT_VALUES_FILE, BENCHMARKS_FILE):
        if not path.exists():
            raise FileNotFoundError(
                f"Missing {path.name}. Download the official workbook into {DOWNLOAD_DIR}."
            )

    return {
        "schema_version": 1,
        "catalog_version": "2026-08-10",
        "generated_on": date.today().isoformat(),
        "scope": "Türkiye origin defaults and definitive-period CBAM benchmarks",
        "sources": [
            {
                "id": "eu-default-values-2026-08-10",
                "title": "Default values definitive period, corrected workbook",
                "publisher": "European Commission, DG TAXUD",
                "published_on": "2026-08-10",
                "legal_basis": "Commission Implementing Regulation (EU) 2026/1740 correcting (EU) 2025/2621",
                "url": DEFAULT_VALUES_URL,
                "sha256": sha256(DEFAULT_VALUES_FILE),
                "binding": False,
            },
            {
                "id": "eu-benchmarks-2026-02-13",
                "title": "Benchmarks definitive period workbook",
                "publisher": "European Commission, DG TAXUD",
                "published_on": "2026-02-13",
                "legal_basis": "Commission Implementing Regulation (EU) 2025/2620",
                "url": BENCHMARKS_URL,
                "sha256": sha256(BENCHMARKS_FILE),
                "binding": False,
            },
            {
                "id": "eu-cbam-prices-2026",
                "title": "Price of CBAM certificates",
                "publisher": "European Commission, DG TAXUD",
                "published_on": "2026-07-06",
                "legal_basis": "Commission Implementing Regulation (EU) 2025/2548",
                "url": "https://taxation-customs.ec.europa.eu/carbon-border-adjustment-mechanism/price-cbam-certificates_en",
                "binding": True,
            },
            {
                "id": "eu-free-allocation-guidance-2026-08-14",
                "title": "Guidance 4, free allocation adjustment",
                "publisher": "European Commission, DG TAXUD",
                "published_on": "2026-08-14",
                "legal_basis": "Explanatory guidance for Implementing Regulation (EU) 2025/2620",
                "url": "https://taxation-customs.ec.europa.eu/document/download/3aa2c730-4f1b-4524-8f71-d9cfe4880ac7_en",
                "binding": False,
            },
        ],
        "certificate_prices": [
            {
                "period": "2026-Q1",
                "eur_per_tco2e": 75.36,
                "published_on": "2026-04-07",
                "source_id": "eu-cbam-prices-2026",
            },
            {
                "period": "2026-Q2",
                "eur_per_tco2e": 75.28,
                "published_on": "2026-07-06",
                "source_id": "eu-cbam-prices-2026",
            },
        ],
        "cbam_factors": {
            "2026": 0.975,
            "2027": 0.95,
            "2028": 0.9,
            "2029": 0.775,
            "2030": 0.515,
            "2031": 0.39,
            "2032": 0.265,
            "2033": 0.14,
            "2034": 0.0,
        },
        "cscf": {
            "2026": {
                "value": 1.0,
                "status": "preliminary",
                "source_id": "eu-free-allocation-guidance-2026-08-14",
            }
        },
        "threshold": {
            "annual_tonnes": 50.0,
            "aggregation_scope": "EU importer aggregate for cement, fertilisers, iron and steel, and aluminium",
            "legal_basis": "Regulation (EU) 2025/2083, Article 2a and Annex VII",
        },
        "default_values": parse_default_values(DEFAULT_VALUES_FILE),
        "benchmarks": parse_benchmarks(BENCHMARKS_FILE),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Rebuild the versioned CBAM reference catalog."
    )
    parser.add_argument(
        "--download",
        action="store_true",
        help="Download the pinned official workbooks before rebuilding.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.download:
        for url, destination in (
            (DEFAULT_VALUES_URL, DEFAULT_VALUES_FILE),
            (BENCHMARKS_URL, BENCHMARKS_FILE),
        ):
            print(f"Downloading {destination.name} from {urlparse(url).hostname}...")
            download_file(url, destination)

    catalog = build_catalog()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        f"Wrote {OUTPUT} with {len(catalog['default_values'])} Türkiye defaults "
        f"and {len(catalog['benchmarks'])} benchmark rows."
    )


if __name__ == "__main__":
    main()
